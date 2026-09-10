"""Regenerate data/MISCELLANEOUS_TECH_TIPS.xlsx from the canonical
data/tech_tips.csv. The CSV is the source of truth (it diffs cleanly in git);
the xlsx is a generated deliverable formatted for the VAC team to hand out.

Usage: python scripts/build_xlsx.py [csv_path] [xlsx_path]
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_CSV = REPO_ROOT / "data" / "tech_tips.csv"
DEFAULT_XLSX = REPO_ROOT / "data" / "MISCELLANEOUS_TECH_TIPS.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
DEAD_FILL = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
FLAG_FILL = PatternFill(start_color="FFF4CE", end_color="FFF4CE", fill_type="solid")
CANDIDATE_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

DISPLAY_COLUMNS = [
    ("Topic", "topic", 42),
    ("Synopsis", "synopsis", 70),
    ("Category", "category", 20),
    ("Status", "status", 16),
    ("Last Verified", "last_verified", 14),
    ("Notes", "notes", 40),
]


def load_rows(csv_path: Path) -> list[dict]:
    with csv_path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def build(csv_path: Path = DEFAULT_CSV, xlsx_path: Path = DEFAULT_XLSX) -> Path:
    rows = load_rows(csv_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "MISC TECH TIPS"

    for col_idx, (label, _key, width) in enumerate(DISPLAY_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, (_label, key, _width) in enumerate(DISPLAY_COLUMNS, start=1):
            value = row.get(key, "")
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        topic_cell = ws.cell(row=r_idx, column=1)
        if row.get("link"):
            topic_cell.hyperlink = row["link"]
            topic_cell.font = LINK_FONT

        status = (row.get("status") or "").lower()
        fill = None
        if status == "dead_link":
            fill = DEAD_FILL
        elif status in ("possibly_superseded", "stale"):
            fill = FLAG_FILL
        elif status == "candidate_new":
            fill = CANDIDATE_FILL
        if fill:
            for c_idx in range(1, len(DISPLAY_COLUMNS) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = fill

    legend_row = len(rows) + 3
    ws.cell(row=legend_row, column=1, value="Legend:").font = Font(name="Arial", bold=True, size=9)
    legend = [
        (DEAD_FILL, "Dead link — flagged this run, PR pending review"),
        (FLAG_FILL, "Possibly superseded / stale — verify before next agent training"),
        (CANDIDATE_FILL, "New candidate tip discovered — not yet confirmed"),
    ]
    for i, (fill, text) in enumerate(legend, start=1):
        c = ws.cell(row=legend_row + i, column=1, value=text)
        c.font = Font(name="Arial", size=9, italic=True)
        c.fill = fill

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    return xlsx_path


if __name__ == "__main__":
    csv_arg = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_CSV
    xlsx_arg = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_XLSX
    out = build(csv_arg, xlsx_arg)
    print(f"Wrote {out}")
